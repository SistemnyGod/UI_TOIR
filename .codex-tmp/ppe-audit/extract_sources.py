import json
import sys
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        return " ".join(value.replace("\xa0", " ").split())
    return value


def read_workbook(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        rows = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [clean(value) for value in row]
            if any(value not in (None, "") for value in values):
                rows.append({"row": row_index, "values": values})
        sheets.append({
            "title": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "rows": rows,
        })
    return {"path": str(path), "sheets": sheets}


def read_document(path):
    document = Document(path)
    paragraphs = [clean(paragraph.text) for paragraph in document.paragraphs if clean(paragraph.text)]
    tables = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row_index, row in enumerate(table.rows, start=1):
            values = [clean(cell.text) for cell in row.cells]
            rows.append({"row": row_index, "values": values})
        tables.append({"table": table_index, "rows": rows})
    return {"path": str(path), "paragraphs": paragraphs, "tables": tables}


def main():
    output = Path(sys.argv[1])
    workbook_paths = [Path(value) for value in sys.argv[2:4]]
    document_paths = [Path(value) for value in sys.argv[4:]]
    payload = {
        "workbooks": [read_workbook(path) for path in workbook_paths],
        "documents": [read_document(path) for path in document_paths],
    }
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "output": str(output),
        "workbooks": [
            {"path": item["path"], "sheets": [{"title": sheet["title"], "rows": len(sheet["rows"]), "max_row": sheet["max_row"], "max_column": sheet["max_column"]} for sheet in item["sheets"]]}
            for item in payload["workbooks"]
        ],
        "documents": [
            {"path": item["path"], "paragraphs": len(item["paragraphs"]), "tables": len(item["tables"]), "table_rows": sum(len(table["rows"]) for table in item["tables"])}
            for item in payload["documents"]
        ],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
