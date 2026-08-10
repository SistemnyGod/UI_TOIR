from openpyxl import load_workbook
from pathlib import Path
import json
import re

files = [
    Path(r"C:\Users\AI_server\Documents\Нормы выдачи СИЗ АТОМ (актуальная) 01.04.26.xlsx"),
    Path(r"C:\Users\AI_server\Documents\спецодежда номенклатура.xlsx"),
]
pattern = re.compile(r"обув|ботин|сапог|зимн", re.I)

for path in files:
    wb = load_workbook(path, read_only=True, data_only=True)
    matches = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if pattern.search(" | ".join(values)):
                matches.append({"sheet": ws.title, "row": row[0] if row else None, "values": values[:15]})
                if len(matches) >= 80:
                    break
        if len(matches) >= 80:
            break
    print(json.dumps({"file": str(path), "sheets": wb.sheetnames, "matches": matches}, ensure_ascii=False))
